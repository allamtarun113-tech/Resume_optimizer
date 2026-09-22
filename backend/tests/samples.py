"""Sample documents and recorded LLM outputs used across tests."""

import io

from docx import Document

from app.schemas.profile import Certification, Education, Project, SkillMention, StudentProfile
from app.schemas.requirements import JobRequirements, Requirement

RESUME_CLASSIC = """Priya Sharma
priya.sharma@example.com | +91 98765 43210 | github.com/priya

SUMMARY
Final-year CSE student who likes building backend systems.

EDUCATION
B.Tech in Computer Science, ABC Institute of Technology    2022 - 2026
CGPA: 8.7/10

TECHNICAL SKILLS
Languages: Python, Java, SQL
Frameworks: FastAPI, React
Tools: Git, Docker, PostgreSQL

PROJECTS
Campus Food Ordering App
• Built a FastAPI backend with PostgreSQL serving 2,000 students
• Added JWT auth and Redis caching, cutting p95 latency from 800ms to 120ms
Stock Price Predictor
• Trained an LSTM in PyTorch on 5 years of NSE data

EXPERIENCE
Backend Intern, Acme Payments    May 2025 - Jul 2025
• Wrote REST endpoints in Go for refund processing

CERTIFICATIONS
AWS Certified Cloud Practitioner (2024)
"""

RESUME_COLON_HEADINGS = """Rahul Verma
Objective:
Aspiring data analyst.
Education:
M.Sc Statistics, XYZ University, 2025
Skills:
Excel, Tableau, Python (pandas), SQL
Academic Projects:
Sales dashboard in Tableau for a retail chain
Achievements:
Winner, college hackathon 2024
"""

SUPPORTING_PROJECT = """Kubernetes deployment for the Campus Food Ordering App.
I containerized the FastAPI service with Docker and deployed it on a 3-node k3s cluster
with a Horizontal Pod Autoscaler. Set up GitHub Actions CI and Prometheus metrics."""

JOB_DESCRIPTION = """Backend Engineer (New Grad) at Example Corp
Requirements:
- Strong Python and FastAPI
- Experience with PostgreSQL and REST API design
- Familiarity with Docker and Kubernetes
Nice to have: AWS, Redis
- Good communication skills
"""


def _escape_pdf(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(text: str) -> bytes:
    """Build a minimal one-page PDF with a real text layer (Helvetica, one line per row)."""
    lines = [line.encode("latin-1", "replace").decode("latin-1") for line in text.split("\n")]
    ops = ["BT", "/F1 10 Tf", "12 TL", "40 800 Td"]
    ops += [f"({_escape_pdf(line)}) Tj T*" for line in lines]
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\n" % (len(objects) + 1))
    out.write(b"startxref\n%d\n%%%%EOF\n" % xref)
    return out.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# Recorded LLM outputs. Sources use the prompt labels (RESUME, S1, ...), as the model returns.
PROFILE_FIXTURE = StudentProfile(
    headline="Final-year CSE student who likes building backend systems.",
    skills=[
        SkillMention(
            name="Python",
            context="skills_section",
            context_name=None,
            source="RESUME",
            evidence="Languages: Python, Java, SQL",
        ),
        SkillMention(
            name="FastAPI",
            context="project",
            context_name="Campus Food Ordering App",
            source="RESUME",
            evidence="Built a FastAPI backend with PostgreSQL serving 2,000 students",
        ),
        SkillMention(
            name="Kubernetes",
            context="project",
            context_name="Campus Food Ordering App",
            source="S1",
            evidence="deployed it on a 3-node k3s cluster",
        ),
        SkillMention(
            name="Made-up skill",
            context="other",
            context_name=None,
            source="S9",  # not a real document: must be dropped
            evidence="n/a",
        ),
    ],
    projects=[
        Project(
            name="Campus Food Ordering App",
            summary="Food ordering backend for a college campus.",
            role=None,
            technologies=["FastAPI", "PostgreSQL", "Redis", "JWT"],
            highlights=["Added JWT auth and Redis caching"],
            metrics=["2,000 students", "p95 latency 800ms -> 120ms"],
            challenges=[],
            start=None,
            end=None,
            url=None,
            source="RESUME",
            evidence="Built a FastAPI backend with PostgreSQL serving 2,000 students",
        )
    ],
    experience=[],
    education=[
        Education(
            institution="ABC Institute of Technology",
            degree="B.Tech",
            field_of_study="Computer Science",
            start="2022",
            end="2026",
            grade="8.7/10",
            source="RESUME",
            evidence="B.Tech in Computer Science, ABC Institute of Technology",
        )
    ],
    certifications=[
        Certification(
            name="AWS Certified Cloud Practitioner",
            issuer="AWS",
            date="2024",
            source="RESUME",
            evidence="AWS Certified Cloud Practitioner (2024)",
        )
    ],
)


def _req(name: str, category: str, importance: str, evidence: str) -> Requirement:
    return Requirement.model_validate(
        {
            "name": name,
            "category": category,
            "importance": importance,
            "min_years": None,
            "evidence": evidence,
        }
    )


REQUIREMENTS_FIXTURE = JobRequirements(
    role_title="Backend Engineer",
    company="Example Corp",
    seniority="entry",
    requirements=[
        _req("Python", "skill", "must", "Strong Python and FastAPI"),
        _req("FastAPI", "skill", "must", "Strong Python and FastAPI"),
        _req("Docker", "skill", "nice", "Familiarity with Docker and Kubernetes"),
        _req("docker", "skill", "must", "Familiarity with Docker"),  # duplicate, stricter
        _req("AWS", "skill", "nice", "Nice to have: AWS, Redis"),
        _req("Communication", "soft_skill", "must", "Good communication skills"),
    ],
)
