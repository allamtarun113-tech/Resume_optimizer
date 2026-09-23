"""ATS check: layout facts, text facts, keyword facts, the pure scoring rules, and the
pipeline/API wiring."""

import io
from collections.abc import Iterator
from fractions import Fraction
from typing import Any

import pytest
from docx import Document

from app.agents.ats_checker import AtsChecker, keyword_facts, text_facts
from app.agents.skill_normalizer import SkillNormalizer
from app.core.config import Settings
from app.parsing.layout import extract_layout
from app.schemas.ats import KeywordFact, ResumeLayout, TextFacts
from app.scoring.ats import (
    ats_report,
    ats_score,
    final_score,
    keyword_check,
    keyword_share,
    layout_checks,
    readable_text,
    section_checks,
)
from app.scoring.weights import ATS_POINTS, ATS_VERSION
from app.skills.taxonomy import load_taxonomy
from tests.builders import profile, req, reqs, skill
from tests.samples import RESUME_CLASSIC, make_docx, make_pdf
from tests.test_api_analyses import Harness
from tests.test_explain import ANSWER, explain, explainer_inputs, finished

TAXONOMY = load_taxonomy()

CLEAN = ResumeLayout(
    file_type="pdf",
    pages=1,
    tables=0,
    multi_column=False,
    images=0,
    photo=False,
    text_boxes=0,
    header_footer_contact=None,
)


def facts(**over: Any) -> TextFacts:
    base: dict[str, Any] = {
        "chars": 2000,
        "unreadable_chars": 0,
        "words": 400,
        "sections": ["education", "experience", "skills"],
        "has_email": True,
        "has_phone": True,
        "link_urls": ["linkedin.com/in/asha"],
        "link_mentions": [],
        "date_styles": {"Month YYYY": "May 2025"},
        "personal_details": [],
    }
    return TextFacts(**{**base, **over})


def kw(i: int, found: bool, importance: str = "must") -> KeywordFact:
    return KeywordFact(
        requirement_index=i,
        name=f"Skill {i}",
        importance=importance,
        found=found,
        bucket=None,
    )


def by_id(checks: list[Any]) -> dict[str, Any]:
    return {c.id: c for c in checks}


# -- layout --------------------------------------------------------------------------------


def test_simple_pdf_has_a_clean_layout() -> None:
    layout = extract_layout(make_pdf(RESUME_CLASSIC), "resume.pdf")
    assert layout == CLEAN


def test_two_columns_are_detected_but_right_aligned_dates_are_not() -> None:
    left = "\n".join(f"Built a service number {i} in Python" for i in range(10))
    right = "\n".join(f"Led the robotics club team {i}" for i in range(10))
    two = extract_layout(make_pdf(left, right), "r.pdf")
    assert two is not None and two.multi_column

    dates = "\n".join("May 2023 to Jul 2023" for _ in range(10))
    one = extract_layout(make_pdf(left, dates), "r.pdf")
    assert one is not None and not one.multi_column


def test_docx_layout_finds_tables_columns_and_header_contact() -> None:
    doc = Document()
    doc.add_paragraph("Asha Rao")
    doc.sections[0].header.paragraphs[0].text = "asha@example.com | +91 98765 43210"
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "B.Tech"
    table.cell(1, 1).text = "2026"
    cols = doc.sections[0]._sectPr.xpath("./w:cols")
    if cols:
        cols[0].set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num", "2")
    else:  # pragma: no cover - python-docx templates always have w:cols
        pytest.fail("template without w:cols")
    buffer = io.BytesIO()
    doc.save(buffer)
    layout = extract_layout(buffer.getvalue(), "resume.docx")
    assert layout is not None
    assert (layout.file_type, layout.tables, layout.pages) == ("docx", 1, None)
    assert layout.multi_column and layout.header_footer_contact

    plain = extract_layout(make_docx(["Asha Rao", "asha@example.com"]), "resume.docx")
    assert plain is not None and not plain.multi_column and not plain.header_footer_contact


def test_unreadable_files_have_no_layout() -> None:
    assert extract_layout(b"plain text resume", "resume.txt") is None
    assert extract_layout(b"%PDF-1.4 broken", "resume.pdf") is None


# -- text facts ----------------------------------------------------------------------------


def test_text_facts_on_a_resume_with_common_problems() -> None:
    text = (
        "Akhila M\nFemale, 21\n9398507060\nakhila@example.com\nLinkedIn\n"
        "EDUCATION\nB.Tech 2020 to 2024\nTECHNICAL SKILLS\nPython\n"
        "PROJECTS\nBus booking May 2023 to 07/2023\nMarital status: single\n\ufffd\ue000"
    )
    f = text_facts(text)
    assert f.has_email and f.has_phone
    assert f.link_urls == [] and f.link_mentions == ["LinkedIn"]
    assert f.date_styles == {"Month YYYY": "May 2023", "MM/YYYY": "07/2023"}
    assert f.personal_details == ["gender", "age", "marital status"]
    assert f.sections == ["education", "projects", "skills"]
    assert f.unreadable_chars == 2


def test_text_facts_on_a_clean_resume() -> None:
    f = text_facts(RESUME_CLASSIC + "\nlinkedin.com/in/priya-s. GitHub")
    assert f.link_urls == ["github.com/priya", "linkedin.com/in/priya-s"]
    assert f.link_mentions == []
    assert f.personal_details == []
    assert f.has_email and f.has_phone
    assert not text_facts("Call 2020-2024 or 12345").has_phone
    assert text_facts("x (cid:12) y").unreadable_chars == 5


def test_keywords_are_matched_word_for_word_or_by_alias() -> None:
    requirements = reqs(
        req("PostgreSQL"),
        req("SQL"),  # only implied by PostgreSQL: an ATS wants the word itself
        req("Kubernetes", importance="nice"),
        req("Payments domain knowledge", "domain"),
        req("Fintech", "domain"),
        req("Communication", "soft_skill"),  # not a keyword check
        req("Java or Kotlin"),
    )
    p = profile(skills=[skill("PostgreSQL")])
    normalized = SkillNormalizer(TAXONOMY).run(p, requirements)
    text = "Built APIs on Postgres for a fintech startup in Kotlin."
    found = {k.name: k.found for k in keyword_facts(requirements, normalized, [], text, TAXONOMY)}
    assert found == {
        "PostgreSQL": True,
        "SQL": False,
        "Kubernetes": False,
        "Payments domain knowledge": False,
        "Fintech": True,
        "Java or Kotlin": True,
    }


# -- scoring rules -------------------------------------------------------------------------


def test_clean_resume_scores_full_marks() -> None:
    report = ats_report(CLEAN, facts(), [kw(0, True), kw(1, True, "nice")])
    assert report.score == report.potential_score == 100
    assert report.version == ATS_VERSION
    assert all(c.status == "pass" and c.fix is None for c in report.checks)
    assert {c.id for c in report.checks} == set(ATS_POINTS) - {"contact_in_body"}


def test_readable_text_levels() -> None:
    assert readable_text(facts(unreadable_chars=2)).status == "pass"
    warn = readable_text(facts(unreadable_chars=20))
    assert (warn.status, warn.points) == ("warn", ATS_POINTS["readable_text"] / 2)
    assert warn.fix
    assert readable_text(facts(unreadable_chars=200)).status == "fail"


def test_layout_problems_lose_points() -> None:
    bad = CLEAN.model_copy(
        update={
            "multi_column": True,
            "tables": 2,
            "images": 3,
            "photo": True,
            "text_boxes": 1,
            "pages": 3,
            "header_footer_contact": True,
        }
    )
    checks = by_id(layout_checks(bad))
    assert checks["single_column"].status == "fail"
    assert checks["no_tables"].status == "warn" and "2 tables" in checks["no_tables"].detail
    assert checks["no_graphics"].items == ["a photo", "2 pictures or icons", "1 text box"]
    assert checks["no_graphics"].points == 0
    assert checks["contact_in_body"].status == "fail"
    assert checks["page_count"].status == "warn"

    icons = by_id(
        layout_checks(CLEAN.model_copy(update={"images": 1, "header_footer_contact": False}))
    )
    assert icons["no_graphics"].status == "warn"
    assert icons["no_graphics"].items == ["1 picture or icon"]
    assert icons["contact_in_body"].status == "pass"
    assert "1 page" in icons["page_count"].detail

    docx = CLEAN.model_copy(update={"pages": None})
    assert "page_count" not in by_id(layout_checks(docx))


def test_section_problems_lose_points() -> None:
    checks = by_id(
        section_checks(
            facts(
                sections=["projects"],
                has_phone=False,
                has_email=False,
                link_urls=[],
                link_mentions=["LinkedIn", "GitHub"],
                date_styles={"Month YYYY": "May 2023", "MM/YYYY": "07/2023"},
                personal_details=["gender"],
                words=120,
            )
        )
    )
    assert checks["standard_headings"].items == ["Education", "Skills"]
    assert checks["contact_details"].status == "fail"
    assert checks["contact_details"].items == ["email", "phone number"]
    assert checks["profile_links"].status == "warn"
    assert "LinkedIn and GitHub" in checks["profile_links"].detail
    assert checks["date_format"].items == ["May 2023", "07/2023"]
    assert checks["personal_details"].status == "warn"
    assert checks["personal_details"].points == 0
    assert "short" in checks["length"].detail

    other = by_id(section_checks(facts(has_phone=False, link_urls=[], words=1500)))
    assert other["contact_details"].status == "warn"
    assert "no LinkedIn" in other["profile_links"].detail
    assert "long" in other["length"].detail


def test_keyword_share_counts_must_haves_three_times() -> None:
    keywords = [kw(0, True), kw(1, False, "nice"), kw(2, False)]
    assert keyword_share(keywords) == Fraction(3, 7)
    assert keyword_share(keywords, {2}) == Fraction(6, 7)
    assert keyword_share([]) == 1

    assert keyword_check(keywords).status == "warn"
    assert keyword_check(keywords).items == ["Skill 1", "Skill 2"]
    assert keyword_check(keywords, {1, 2}).status == "pass"
    assert keyword_check([kw(0, False)]).status == "fail"
    assert "no specific skills" in keyword_check([]).detail


def test_potential_score_counts_the_suggested_keywords() -> None:
    report = ats_report(None, facts(), [kw(0, True), kw(1, False)], addressed={1})
    assert report.score < report.potential_score == 100
    # Without a file, the layout checks don't apply (and don't count).
    assert "single_column" not in {c.id for c in report.checks}


def test_scores_and_final_score() -> None:
    checks = section_checks(facts(personal_details=["age"]))
    total = sum(ATS_POINTS[c.id] for c in checks)
    assert ats_score(checks) == round(100 * (total - ATS_POINTS["personal_details"]) / total)
    assert final_score(35, 72) == 46  # 0.7 x 35 + 0.3 x 72 = 46.1
    assert final_score(100, 100) == 100
    assert final_score(0, 0) == 0


def test_adding_a_keyword_never_lowers_the_ats_score() -> None:
    keywords = [kw(i, i % 2 == 0, "must" if i % 3 else "nice") for i in range(7)]
    base = ats_report(CLEAN, facts(), keywords).score
    for i in range(7):
        assert ats_report(CLEAN, facts(), keywords, addressed={i}).score >= base


# -- agent + pipeline + API ----------------------------------------------------------------


def test_ats_checker_combines_layout_text_and_keywords() -> None:
    requirements = reqs(req("Python"), req("Rust"))
    normalized = SkillNormalizer(TAXONOMY).run(profile(), requirements)
    report = AtsChecker(TAXONOMY).run(
        layout=CLEAN,
        resume_text=RESUME_CLASSIC,
        requirements=requirements,
        normalized=normalized,
        matches=[],
    )
    assert [(k.name, k.found) for k in report.keywords] == [("Python", True), ("Rust", False)]
    assert by_id(report.checks)["keywords"].items == ["Rust"]


@pytest.fixture
def h(settings: Settings) -> Iterator[Harness]:
    harness = Harness(settings)
    harness.model.responses["Explanation"] = ANSWER
    with harness.client:
        yield harness


def test_analysis_stores_ats_and_final_scores(h: Harness) -> None:
    body = finished(h)
    ats = body["ats"]
    assert body["ats_score"] == ats["score"]
    assert body["final_score"] == final_score(body["fit_score"], body["ats_score"])
    assert body["potential_final_score"] == final_score(
        body["potential_score"], ats["potential_score"]
    )
    checks = by_id([type("C", (), c) for c in ats["checks"]])
    assert "single_column" in checks  # the resume file's layout was read
    assert {k["name"] for k in ats["keywords"]} >= {"Python", "Docker"}

    summary = h.client.get("/analyses").json()[0]
    assert (summary["ats_score"], summary["final_score"]) == (
        body["ats_score"],
        body["final_score"],
    )

    report = h.client.get(f"/analyses/{body['id']}/export").text
    assert f"**Final score:** {body['final_score']}%" in report
    assert f"## ATS check ({body['ats_score']}%)" in report


def test_missing_resume_file_skips_only_the_layout_checks(h: Harness) -> None:
    resume_id = h.upload_resume().json()["id"]
    h.repo.files.clear()
    created = h.analyze(resume_id)
    body = h.client.get(f"/analyses/{created.json()['id']}").json()
    assert body["status"] == "done"
    ids = {c["id"] for c in body["ats"]["checks"]}
    assert "single_column" not in ids and "keywords" in ids


def test_explain_ats_score_and_a_check(h: Harness) -> None:
    body = finished(h)
    assert explain(h, body["id"], "ats").status_code == 200
    assert explain(h, body["id"], "ats_check", "keywords").status_code == 200
    assert explain(h, body["id"], "ats_check", "nope").status_code == 404
    ats_prompt, check_prompt = explainer_inputs(h)
    assert f"ATS score: {body['ats_score']}%" in ats_prompt
    assert "ATS check: Job keywords" in check_prompt
    assert "Final score = 70% job fit + 30% ATS score." in ats_prompt
