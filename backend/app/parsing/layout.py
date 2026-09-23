"""Resume layout facts for the ATS check: tables, columns, pictures, text boxes.

Pure functions over the file's bytes (no I/O, no LLM). Heuristics are deliberately
conservative: a check only fails when the problem is clear."""

import io
import logging
import re
from typing import Any

import pdfplumber
from docx import Document as load_docx

from app.parsing.extract import FileType, detect_file_type
from app.schemas.ats import ResumeLayout

logger = logging.getLogger(__name__)

MAX_PAGES = 5
# A gap wider than this share of the page splits a line into two chunks.
COLUMN_GAP = 0.08
WORD_GAP = 1.5  # points between characters that mean a new word
# A line counts as two columns side by side when both chunks have this many words.
COLUMN_MIN_WORDS = 3
# Multi-column when at least this many (and this share of) lines are side by side.
COLUMN_MIN_LINES = 5
COLUMN_MIN_SHARE = 0.2
# A picture at least this big (points) with a near-square shape on page 1 is a photo.
PHOTO_MIN_SIDE = 40
PHOTO_MAX_RATIO = 1.6
TABLE_MIN_ROWS = 2
TABLE_MIN_COLS = 2

_YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b|\bpresent\b", re.IGNORECASE)
CONTACT_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+|\+?\d[\d\s().-]{8,}\d")


def extract_layout(data: bytes, filename: str) -> ResumeLayout | None:
    """Layout facts, or None for plain text or a file that can't be read."""
    try:
        file_type = detect_file_type(data, filename)
        if file_type is FileType.PDF:
            return _pdf_layout(data)
        if file_type is FileType.DOCX:
            return _docx_layout(data)
    except Exception:  # odd files: the ATS check just skips the layout part
        logger.warning("Could not read the resume layout", exc_info=True)
    return None


def _side_by_side(line: dict[str, Any], width: float, tables: list[tuple[float, float]]) -> bool:
    """A text line made of two chunks of prose far apart (two columns), outside tables.
    A chunk with a year is a right-aligned date ("Intern ... May 2023"), not a column."""
    if any(top <= line["top"] <= bottom for top, bottom in tables):
        return False
    chars = sorted(line["chars"], key=lambda c: c["x0"])
    chunks: list[str] = []
    current: list[str] = []
    last_x1: float | None = None
    for c in chars:
        gap = c["x0"] - last_x1 if last_x1 is not None else 0
        if gap > width * COLUMN_GAP:
            chunks.append("".join(current))
            current = []
        elif gap > WORD_GAP:  # some PDFs have no space characters between words
            current.append(" ")
        current.append(c["text"])
        last_x1 = c["x1"]
    chunks.append("".join(current))
    prose = [ch for ch in chunks if len(ch.split()) >= COLUMN_MIN_WORDS and not _YEAR_RE.search(ch)]
    return len(chunks) >= 2 and len(prose) >= 2


def _pdf_layout(data: bytes) -> ResumeLayout:
    tables = images = side_by_side = lines = 0
    photo = False
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = len(pdf.pages)
        for number, page in enumerate(pdf.pages[:MAX_PAGES]):
            found = [
                t
                for t in page.find_tables()
                if len(t.rows) >= TABLE_MIN_ROWS
                and max(len(r.cells) for r in t.rows) >= TABLE_MIN_COLS
            ]
            tables += len(found)
            spans = [(t.bbox[1], t.bbox[3]) for t in found]
            for image in page.images:
                images += 1
                w, h = float(image["width"]), float(image["height"])
                if (
                    number == 0
                    and min(w, h) >= PHOTO_MIN_SIDE
                    and max(w, h) / min(w, h) <= PHOTO_MAX_RATIO
                    and float(image["top"]) < page.height / 3
                ):
                    photo = True
            for line in page.extract_text_lines():
                lines += 1
                side_by_side += _side_by_side(line, float(page.width), spans)
    multi = side_by_side >= COLUMN_MIN_LINES and side_by_side >= lines * COLUMN_MIN_SHARE
    return ResumeLayout(
        file_type="pdf",
        pages=pages,
        tables=tables,
        multi_column=multi,
        images=images,
        photo=photo,
        text_boxes=0,
        header_footer_contact=None,
    )


def _docx_layout(data: bytes) -> ResumeLayout:
    doc = load_docx(io.BytesIO(data))
    body = doc.element.body.xml
    columns = re.findall(r'<w:cols\b[^>]*\bw:num="(\d+)"', body)
    header_text = " ".join(
        p.text
        for section in doc.sections
        for part in (section.header, section.footer)
        for p in part.paragraphs
    )
    body_text = " ".join(p.text for p in doc.paragraphs)
    in_header = bool(CONTACT_RE.search(header_text)) and not CONTACT_RE.search(body_text)
    tables = [
        t for t in doc.tables if len(t.rows) >= TABLE_MIN_ROWS and len(t.columns) >= TABLE_MIN_COLS
    ]
    return ResumeLayout(
        file_type="docx",
        pages=None,
        tables=len(tables),
        multi_column=any(int(n) > 1 for n in columns),
        images=len(doc.inline_shapes) + body.count("<wp:anchor"),
        photo=False,
        text_boxes=body.count("<w:txbxContent"),
        header_footer_contact=in_header,
    )
